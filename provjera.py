# -*- coding: utf-8 -*-
"""Samostalna provjera aplikacije (pokreće se iz .exe: Sihtarica_IDSS_IMH.exe --provjera).
Generiše vlastite testne izvoze uređaja i provjerava sve scenarije. Ne dira stvarnu evidenciju."""
import datetime as dt, json, os, sys, tempfile, traceback, shutil
from pathlib import Path
import openpyxl
import engine as E, izvoz as I

def resurs(ime): return Path(getattr(sys, "_MEIPASS", Path(__file__).parent)) / ime

IMENA = [("MK001", "Amar Hadžić", "Uprava"), ("MK002", "Lejla Kovačević", "Nastava"), ("MK003", "Đorđe Šćepanović", "Nastava"),
         ("MK008", "Mia Alić", "IMH – odgajatelji"), ("MK019", "Nermin Imamović", "IMH – odgajatelji")]

def tabelarni(put, g, m, od, do, posebno=None):
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Attendance Logs"
    ws["A1"] = "Attendance Logs"; ws["A2"] = "Date"; ws["C2"] = f"{od:02d}/{m:02d}/{g} ~ {do:02d}/{m:02d}/{g}"
    ws.append([]); ws.append(["Date", "Employee ID", "Employee name", "Department", "Clock-in", "Clock-out", "Break (min)",
                              "Worked hours", "Status", "Overtime (hours)", "Notes"])
    for d in range(od, do + 1):
        datum = dt.date(g, m, d)
        for rid, ime, odj in IMENA:
            vik = datum.weekday() >= 5
            red = [dt.datetime(g, m, d), rid, ime, odj, None if vik else dt.time(8, 0), None if vik else dt.time(16, 0), 0 if vik else 30,
                   0 if vik else 8.0, "Vikend / neradni dan" if vik else "Prisutan", 0, "TEST"]
            if posebno: red = posebno(datum, rid, red)
            ws.append(red)
    for n in ("Schedule information sheet", "Attendance summary", "Abnormal"): wb.create_sheet(n)
    wb.save(put)

def mrezni(put, g, m):
    n = E.zadnji_dan(g, m)
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Attendance Logs"
    ws["A1"] = "Attendance Logs"; ws["A2"] = "Date"; ws["C2"] = f"01/{m:02d}/{g} ~ {n:02d}/{m:02d}/{g}"
    for d in range(1, n + 1): ws.cell(3, d, d)
    r = 4
    for rid, ime, odj in IMENA[:2]:
        ws.cell(r, 1, "ID:"); ws.cell(r, 3, rid); ws.cell(r, 9, "Name:"); ws.cell(r, 11, ime); ws.cell(r, 19, "Dept.:"); ws.cell(r, 21, odj)
        for d in range(1, n + 1):
            if dt.date(g, m, d).weekday() < 5: ws.cell(r + 1, d, "07:55\n16:02")
        r += 2
    ws.cell(5, 2).value = "07:58"; ws.cell(5, 3).value = "07:50\n12:00\n12:30\n16:05"; ws.cell(5, 4).value = "07:50\n10:00\n16:05"
    ws.cell(7, 4).value = "08:01\n08:02\n15:59"; ws.cell(7, 7).value = None
    wb.save(put)

def zakljucaj(put):
    """Simulira Excel: drži fajl otvorenim bez dijeljenja (Windows)."""
    if os.name != "nt": return None
    import ctypes
    from ctypes import wintypes
    k32 = ctypes.WinDLL("kernel32", use_last_error=True)
    k32.CreateFileW.restype = wintypes.HANDLE
    h = k32.CreateFileW(str(put), 0x80000000 | 0x40000000, 0, None, 3, 0x80, None)
    return (k32, h) if h not in (None, wintypes.HANDLE(-1).value) else None

def pokreni(izlaz_folder=None):
    if izlaz_folder: Path(izlaz_folder).mkdir(parents=True, exist_ok=True)
    rez, baza = [], Path(tempfile.mkdtemp(prefix="Sihtarica provjera čćšđž "))
    def test(naziv, fn):
        try: fn(); rez.append((naziv, True, "")); print("  ✔", naziv)
        except Exception as e: rez.append((naziv, False, f"{e}\n{traceback.format_exc(limit=2)}")); print("  ✘", naziv, "–", e)
    ev = baza / "Evidencija IDSS–IMH čćžšđ.xlsx"
    post = dict(E.PODRAZUMIJEVANE_POSTAVKE)
    def uvoz(src, kat=None, satnica=40):
        p = E.napravi_plan(E.citaj_izvor(src, post), ev, post)
        for r in p["radnici"].values():
            if r["novi"]: r["sedmicno"] = satnica
        for k in p["kategorizacija"]:
            if kat: kat(k)
        return E.upisi(p, ev)
    # izvori
    tabelarni(baza / "okt.xlsx", 2026, 10, 1, 31, lambda d, rid, r: (r.__setitem__(8, "Godišnji odmor") or r.__setitem__(4, None) or r.__setitem__(5, None) or r)
              if (rid == "MK019" and d.day == 22) else (r.__setitem__(8, "Odsustvo") or r.__setitem__(4, None) or r.__setitem__(5, None) or r)
              if (rid == "MK008" and d.day == 14) else (r.__setitem__(5, dt.time(16, 30)) or r.__setitem__(9, 0.5) or r) if (rid == "MK001" and d.day == 8) else r)
    tabelarni(baza / "nov_1-15.xlsx", 2026, 11, 1, 15)
    tabelarni(baza / "nov_1-30.xlsx", 2026, 11, 1, 30, lambda d, rid, r: (r.__setitem__(4, dt.time(7, 30)) or r) if (rid == "MK001" and d.day == 3) else
              (r.__setitem__(5, None) or r) if (rid == "MK002" and d.day == 20) else r)
    mrezni(baza / "dec_mreza.xlsx", 2026, 12)
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Attendance Logs"; ws["C2"] = "01/10/2026 ~ 09/10/2026"; wb.save(baza / "prazan.xlsx")
    (baza / "pokvaren.xlsx").write_bytes(b"ovo nije excel")

    test("Kreiranje evidencije iz propisanog obrasca (putanja s č/ć/š i razmacima)", lambda: E.kreiraj_evidenciju(resurs("sablon_Sihtarica_IDSS_IMH.xlsx"), ev))
    def t_okt():
        r = uvoz(baza / "okt.xlsx", lambda k: k.update(kategorija="Neplaćeno odsustvo") if k["id"] == "MK008" else None)
        assert r["novih"] == 5, r
        wb = openpyxl.load_workbook(ev)
        s = wb["2026-10 IDSS MK001"]; assert s["D7"].value == dt.time(8, 0) and s["E14"].value == dt.time(16, 30) and s["I14"].value == 0.5
        s = wb["2026-10 IMH MK019"]; assert s["L28"].value == 8 and "prema uređaju" in s["P28"].value
        s = wb["2026-10 IMH MK008"]; assert s["P20"].value.startswith("Neplaćeno odsustvo")
        assert wb["2026-10 IDSS MK003"]["F3"].value == "Đorđe Šćepanović"
    test("Oktobar: tabelarni izvoz, prekovremeni, godišnji s uređaja, ručna kategorija, naša slova", t_okt)
    def t_isti():
        try: uvoz(baza / "okt.xlsx"); raise AssertionError("isti izvoz je prihvaćen")
        except E.GreskaUpisa as e: assert "već upisan" in str(e)
    test("Isti izvoz se ne može upisati dvaput", t_isti)
    def t_nov():
        r1 = uvoz(baza / "nov_1-15.xlsx"); assert r1["novih"] == 5
        s = openpyxl.load_workbook(ev)["2026-11 IDSS MK001"]; assert s["P22"].value == E.MARKER_VAN
        r2 = uvoz(baza / "nov_1-30.xlsx"); assert r2["dopunjenih"] == 5 and r2["novih"] == 0, r2
        wb = openpyxl.load_workbook(ev); s = wb["2026-11 IDSS MK001"]
        assert s["D9"].value == dt.time(8, 0), "postojeći dan je prepisan"
        assert s["D22"].value == dt.time(8, 0) and s["P22"].value is None, "dan 16.11. nije dopunjen"
        assert any("zadržano postojeće" in str(r[5]) for r in wb["Upozorenja"].iter_rows(min_row=3, values_only=True))
        assert wb["2026-11 IDSS MK002"]["P26"].value == "Nedostaje prijava odlaska"
    test("Novembar: nepotpun mjesec upisan, zatim dopunjen; postojeći dan nije prepisan; sukob zabilježen", t_nov)
    def t_dec():
        assert uvoz(baza / "dec_mreza.xlsx")["novih"] == 2
        s = openpyxl.load_workbook(ev)["2026-12 IDSS MK001"]
        assert s["D8"].value is None and "Jedna prijava u 07:58" in s["P8"].value
        assert (s["F9"].value, s["G9"].value) == (dt.time(12, 0), dt.time(12, 30))
        assert s["D10"].value == dt.time(7, 50) and "3 prijave" in s["P10"].value and s["F10"].value is None
        s = openpyxl.load_workbook(ev)["2026-12 IDSS MK002"]
        assert "Ponovljena prijava" in s["P10"].value and s["P13"].value == E.MARKER_NEMA
    test("Decembar: prijave po danima – 1, 3, 4 prijave, ponovljena prijava, dan bez prijava", t_dec)
    def t_prazan():
        try: E.citaj_izvor(baza / "prazan.xlsx"); raise AssertionError("prazan izvoz prihvaćen")
        except E.GreskaFormata as e: assert "ne sadrži" in str(e) or "prepoznatom" in str(e)
    test("Prazan izvoz se jasno odbija", t_prazan)
    def t_pokv():
        try: E.citaj_izvor(baza / "pokvaren.xlsx"); raise AssertionError("pokvaren fajl prihvaćen")
        except Exception as e: assert not isinstance(e, AssertionError)
    test("Pokvaren fajl se odbija bez rušenja", t_pokv)
    def t_tamper():
        wb = openpyxl.load_workbook(ev); prije = E._snimak(wb); wb["2026-10 IDSS MK001"]["D7"] = dt.time(9, 0)
        assert E.dokaz_istorije(prije, E._snimak(wb)), "izmjena istorije nije otkrivena"
    test("Dokaz istorije otkriva svaku izmjenu ranije upisane ćelije", t_tamper)
    def t_lock():
        if os.name != "nt": raise RuntimeError("preskočeno – test se izvodi samo na Windowsu")
        tabelarni(baza / "jan.xlsx", 2027, 1, 1, 31)
        prije = E.sha256_fajla(ev); l = zakljucaj(ev); assert l, "zaključavanje nije uspjelo"
        try:
            try: uvoz(baza / "jan.xlsx"); raise AssertionError("upis u zaključan fajl je prošao")
            except E.GreskaUpisa as e: assert "otvoren" in str(e) or "zaključan" in str(e)
        finally: l[0].CloseHandle(l[1])
        assert E.sha256_fajla(ev) == prije, "zaključana evidencija je promijenjena"
    test("Windows: evidencija otvorena u Excelu (zaključana) – upis odbijen, fajl netaknut", t_lock)
    def t_izvoz():
        o = baza / "izvoz"; o.mkdir()
        assert I.izvezi_xlsx(ev, o / "sve.xlsx") is None and (o / "sve.xlsx").stat().st_size > 0
        assert I.izvezi_xlsx(ev, o / "okt_imh.xlsx", "2026-10", "IMH") == 2
        assert I.izvezi_pdf(ev, o / "okt.pdf", "2026-10") == 5 and (o / "okt.pdf").read_bytes()[:5] == b"%PDF-"
        I.izvezi_json(ev, o / "sve.json"); m = json.loads((o / "sve.json").read_text(encoding="utf-8"))
        s = next(x for x in m["sihtarice"] if x["list"] == "2026-10 IDSS MK001"); assert s["ukupno"]["ukupno_sati"] == 176.5 and s["ukupno"]["prekovremeni"] == 0.5
        if izlaz_folder:
            for f in o.iterdir(): E.sigurna_kopija(f, Path(izlaz_folder) / f.name)
            E.sigurna_kopija(ev, Path(izlaz_folder) / ev.name)
    test("Izvoz .xlsx (sve i filtrirano), .pdf (fontovi s č/ć/š), .json (sati = obrazac)", t_izvoz)
    def t_kopije():
        assert len(list((baza / "Sigurnosne_kopije").glob("*.xlsx"))) == 4, "broj sigurnosnih kopija ne odgovara broju upisa"
        d = list(openpyxl.load_workbook(ev)["Dnevnik uvoza"].iter_rows(min_row=3, values_only=True)); assert len(d) == 4, d
    test("Sigurnosna kopija prije svakog upisa i trag u Dnevniku uvoza", t_kopije)
    ok = all(x[1] for x in rez)
    if os.environ.get("GITHUB_ACTIONS"):                    # pale provjere kao vidljive napomene na GitHubu
        for n, o, d in rez:
            if not o: print(f"::error title=PALO: {n}::" + d.strip().replace("%", "%25").replace("\r", "").replace("\n", "%0A")[:3000], flush=True)
    izvj = f"Šihtarica IDSS/IMH v{E.VERZIJA} – provjera {dt.datetime.now():%d.%m.%Y %H:%M}\nSistem: {sys.platform} / Python {sys.version.split()[0]}" \
           f" / {'EXE' if getattr(sys, 'frozen', False) else 'izvorni kod'}\n\n" + \
           "\n".join(f"{'PROŠLO' if o else 'PALO  '}  {n}" + (f"\n        {d}" if d else "") for n, o, d in rez) + \
           f"\n\nUKUPNO: {sum(x[1] for x in rez)}/{len(rez)} – {'SVE PROVJERE SU PROŠLE' if ok else 'IMA GREŠAKA'}\n"
    if izlaz_folder: (Path(izlaz_folder) / "rezultat_provjere.txt").write_text(izvj, encoding="utf-8")
    shutil.rmtree(baza, ignore_errors=True)
    return ok, izvj

if __name__ == "__main__":
    ok, t = pokreni(sys.argv[1] if len(sys.argv) > 1 else None); print(t); sys.exit(0 if ok else 1)
