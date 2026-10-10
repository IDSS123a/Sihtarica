# -*- coding: utf-8 -*-
"""Izvoz evidencije (1) u .xlsx, .json i .pdf. Svi izvozi čitaju isključivo evidencijski fajl – ono što
piše u šihtarici, to piše i u izvozu. Dnevni sati se računaju identično formuli obrasca."""
import datetime as dt, json, re, shutil
from copy import copy
from pathlib import Path
import openpyxl
import engine as E

# ----------------------------------------------------------------------- model
def _vr(v):
    if isinstance(v, dt.datetime): return dt.time(v.hour, v.minute)
    return v if isinstance(v, dt.time) else None

def _broj(v):
    return float(v) if isinstance(v, (int, float)) and not isinstance(v, bool) else None

def procitaj_model(evidencija, filter_period=None, filter_inst=None):
    wb = openpyxl.load_workbook(evidencija)
    reg = E.procitaj_registar(wb)
    periodi = []
    for ws in wb.worksheets:
        m = E.NAZIV_LISTA.match(ws.title)
        if not m: continue
        g, mj, inst, rid = int(m.group(1)), int(m.group(2)), m.group(3), m.group(4)
        if filter_period and f"{g}-{mj:02d}" != filter_period: continue
        if filter_inst and inst != filter_inst: continue
        dani, uk = [], dict(ukupno_sati=0.0, prekovremeni=0.0, terenski=0.0, pripravnost=0.0, godisnji=0.0,
                             bolovanje=0.0, praznik=0.0, sedmicni_odmor=0.0)
        for d in range(1, E.zadnji_dan(g, mj) + 1):
            red = E.PRVI_RED + d - 1; datum = dt.date(g, mj, d)
            D, Ee, F, G = (_vr(ws[f"{k}{red}"].value) for k in "DEFG")
            h = E.sati_obrasca(D, Ee, F, G)
            dan = dict(rb=d, datum=datum.isoformat(), dan=E.DANI[datum.weekday()],
                       pocetak_rada=D.strftime("%H:%M") if D else None, zavrsetak_rada=Ee.strftime("%H:%M") if Ee else None,
                       pocetak_pauze=F.strftime("%H:%M") if F else None, zavrsetak_pauze=G.strftime("%H:%M") if G else None,
                       ukupno_sati=h, prekovremeni=_broj(ws[f"I{red}"].value), terenski=_broj(ws[f"J{red}"].value),
                       pripravnost=_broj(ws[f"K{red}"].value), godisnji=_broj(ws[f"L{red}"].value),
                       bolovanje=_broj(ws[f"M{red}"].value), praznik=_broj(ws[f"N{red}"].value),
                       sedmicni_odmor=_broj(ws[f"O{red}"].value), napomena=ws[f"P{red}"].value or None,
                       potpis=ws[f"Q{red}"].value or None)
            for k in uk:
                if dan[k]: uk[k] += dan[k]
            dani.append(dan)
        uk = {k: E.zaokruzi(v) for k, v in uk.items()}
        r = reg.get(rid, {})
        periodi.append(dict(list=ws.title, period=f"{g}-{mj:02d}", godina=g, mjesec=mj, mjesec_naziv=E.MJESECI[mj - 1],
                            ustanova=inst, radnik=dict(id=rid, ime=ws["F3"].value or r.get("ime"), radno_mjesto=ws["L3"].value,
                            ugovoreno_sedmicno=ws["D4"].value, odgovorno_lice=ws["F4"].value),
                            dani=dani, ukupno=uk))
    periodi.sort(key=lambda p: (p["period"], p["ustanova"], p["radnik"]["ime"] or ""))
    def tabela(naziv, od=3):
        if naziv not in wb.sheetnames: return []
        ws = wb[naziv]; zag = [c.value for c in ws[2]]
        return [{zag[i]: (v.isoformat() if isinstance(v, (dt.date, dt.datetime)) else v) for i, v in enumerate(r) if i < len(zag) and zag[i]}
                for r in ws.iter_rows(min_row=od, values_only=True) if any(x not in (None, "") for x in r)]
    upz = tabela(E.L_UPOZORENJA)
    if filter_period: upz = [u for u in upz if u.get("Period") == filter_period]
    ids = {p["radnik"]["id"] for p in periodi}
    registar = [r for r in reg.values() if not (filter_period or filter_inst) or r["id"] in ids]
    return dict(meta=dict(aplikacija=f"Šihtarica IDSS/IMH v{E.VERZIJA}", izvezeno=dt.datetime.now().isoformat(timespec="seconds"),
                          evidencija=Path(evidencija).name, period=filter_period or "svi", ustanova=filter_inst or "sve",
                          broj_sihtarica=len(periodi)),
                registar=registar, sihtarice=periodi, upozorenja=upz, dnevnik_uvoza=tabela(E.L_DNEVNIK))

def dostupni_periodi(evidencija):
    wb = openpyxl.load_workbook(evidencija, read_only=True)
    p = sorted({f"{m.group(1)}-{m.group(2)}" for n in wb.sheetnames if (m := E.NAZIV_LISTA.match(n))}); wb.close()
    return p

# ----------------------------------------------------------------------- JSON
def izvezi_json(evidencija, cilj, period=None, inst=None):
    model = procitaj_model(evidencija, period, inst)
    Path(cilj).write_text(json.dumps(model, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    return model["meta"]["broj_sihtarica"]

# ----------------------------------------------------------------------- XLSX
def izvezi_xlsx(evidencija, cilj, period=None, inst=None):
    if not period and not inst:
        E.sigurna_kopija(evidencija, cilj); return None
    wb = openpyxl.load_workbook(evidencija)
    zadrzi = []
    for ws in list(wb.worksheets):
        m = E.NAZIV_LISTA.match(ws.title)
        if m:
            ok = (not period or f"{m.group(1)}-{m.group(2)}" == period) and (not inst or m.group(3) == inst)
            if ok: zadrzi.append(ws.title)
            else: wb.remove(ws)
        elif ws.title in E.OBRAZAC.values(): wb.remove(ws)
    wsp = wb[E.L_PREGLED]
    stil6, stil7 = [copy(wsp.cell(6, c)._style) for c in range(1, 8)], [copy(wsp.cell(7, c)._style) for c in range(1, 8)]
    for r in range(6, wsp.max_row + 1):
        for c in range(1, 8): wsp.cell(r, c).value = None
    for i, n in enumerate(sorted(zadrzi, key=lambda n: (n.split(" ")[0], n.split(" ")[1], wb[n]["F3"].value or ""))):
        red = 6 + i
        for c in range(1, 8): wsp.cell(red, c)._style = copy((stil6 if red % 2 == 0 else stil7)[c - 1])
        m = E.NAZIV_LISTA.match(n)
        E._red_pregleda(wsp, n, m.group(3), wb[n]["F3"].value, int(m.group(1)), int(m.group(2)), red=red)
    if E.L_UPOZORENJA in wb.sheetnames and period:
        ws = wb[E.L_UPOZORENJA]
        for r in range(ws.max_row, 2, -1):
            if ws.cell(r, 1).value != period: ws.delete_rows(r)
    if E.L_REGISTAR in wb.sheetnames:
        ws = wb[E.L_REGISTAR]; ids = {E.NAZIV_LISTA.match(n).group(4) for n in zadrzi}
        for r in range(ws.max_row, 2, -1):
            if E._id(ws.cell(r, 1).value) not in ids: ws.delete_rows(r)
    wb.calculation.fullCalcOnLoad = True
    wb.save(cilj)
    return len(zadrzi)

# ----------------------------------------------------------------------- PDF
def _fontovi():
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    import sys
    baza = Path(getattr(sys, "_MEIPASS", Path(__file__).parent)) / "fonts"
    if "Carlito" not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont("Carlito", str(baza / "Carlito-Regular.ttf")))
        pdfmetrics.registerFont(TTFont("Carlito-Bold", str(baza / "Carlito-Bold.ttf")))
        pdfmetrics.registerFont(TTFont("Carlito-Italic", str(baza / "Carlito-Italic.ttf")))
        from reportlab.lib.fonts import addMapping
        addMapping("Carlito", 0, 0, "Carlito"); addMapping("Carlito", 1, 0, "Carlito-Bold"); addMapping("Carlito", 0, 1, "Carlito-Italic")

ZAGLAVLJA = ["R.br.", "Datum", "Dan", "Početak rada", "Završetak rada", "Početak pauze", "Završetak pauze",
             "Ukupno dnevno sati", "Prekovremeni sati", "Terenski rad (sati)", "Pripravnost (sati)", "Godišnji odmor (sati)",
             "Bolovanje (sati)", "Praznik / neradni dan (sati)", "Sedmični odmor (sati)", "Drugo odsustvo / napomena",
             "Potpis / inicijali radnika"]
SIRINE = [6, 12, 12, 12, 12, 12, 12, 14, 13, 13, 13, 13, 13, 15, 14, 28, 18]

def izvezi_pdf(evidencija, cilj, period=None, inst=None):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import (BaseDocTemplate, Frame, PageTemplate, Paragraph, Table, TableStyle, Spacer,
                                    PageBreak, KeepTogether)
    _fontovi()
    model = procitaj_model(evidencija, period, inst)
    C = lambda h: colors.HexColor("#" + h)
    plava, svplava, traka, zuta, amber, tekst = (C(E.BOJA[k]) for k in ("plava", "svijetloplava", "traka", "zuta", "amber", "tekst"))
    st = lambda n, **k: ParagraphStyle(n, fontName=k.pop("font", "Carlito"), **k)
    S = dict(naslov=st("n", font="Carlito-Bold", fontSize=13, textColor=colors.white, alignment=1, leading=16),
             lab=st("l", font="Carlito-Bold", fontSize=8, textColor=plava, leading=10),
             val=st("v", fontSize=9.5, textColor=tekst, leading=11),
             zag=st("z", font="Carlito-Bold", fontSize=6.6, textColor=colors.white, alignment=1, leading=7.6),
             cel=st("c", fontSize=7.4, alignment=1, leading=8.6, textColor=tekst),
             nap=st("p", fontSize=6.4, leading=7.4, textColor=tekst),
             fus=st("f", font="Carlito-Italic", fontSize=6.8, textColor=C("555555"), leading=8.4),
             pot=st("s", fontSize=8, textColor=tekst, leading=10))
    stranica = landscape(A4); marg = 9 * mm
    sirina = stranica[0] - 2 * marg
    kol = [sirina * w / sum(SIRINE) for w in SIRINE]

    def podnozje(cv, doc):
        cv.saveState(); cv.setFont("Carlito", 7); cv.setFillColor(C("777777"))
        cv.drawString(marg, 5 * mm, f"{model['meta']['aplikacija']}  ·  izvor: {model['meta']['evidencija']}  ·  izvezeno {dt.datetime.now():%d.%m.%Y %H:%M}")
        cv.drawRightString(stranica[0] - marg, 5 * mm, f"Stranica {doc.page}")
        cv.restoreState()
    doc = BaseDocTemplate(str(cilj), pagesize=stranica, leftMargin=marg, rightMargin=marg, topMargin=marg, bottomMargin=10 * mm,
                          title="Šihtarice IDSS/IMH", author=model["meta"]["aplikacija"])
    doc.addPageTemplates([PageTemplate("p", [Frame(marg, 10 * mm, sirina, stranica[1] - marg - 10 * mm, 0, 0, 0, 0)], onPage=podnozje)])
    f = lambda v: "" if v in (None, "") else (f"{v:.2f}" if isinstance(v, float) else str(v))
    elementi = []
    # --- naslovna: mjesečni pregled
    nasl = Table([[Paragraph("MJESEČNI PREGLED – IDSS i IMH", S["naslov"])]], colWidths=[sirina], rowHeights=[22])
    nasl.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), plava), ("VALIGN", (0, 0), (-1, -1), "MIDDLE")]))
    elementi += [nasl, Spacer(0, 6),
                 Paragraph(f"Period: <b>{model['meta']['period']}</b> &nbsp;·&nbsp; Ustanova: <b>{model['meta']['ustanova']}</b> "
                           f"&nbsp;·&nbsp; Broj šihtarica: <b>{len(model['sihtarice'])}</b>", S["val"]), Spacer(0, 6)]
    pr = [[Paragraph(x, S["zag"]) for x in ("Ustanova", "Radnik/ca", "Mjesec", "Godina", "Ukupno sati", "Prekovremeni",
                                             "Godišnji", "Bolovanje", "Praznik", "Terenski", "Napomena")]]
    for p in model["sihtarice"]:
        u = p["ukupno"]
        ozn = sum(1 for d in p["dani"] if d["napomena"])
        pr.append([Paragraph(x, S["cel"]) for x in (p["ustanova"], p["radnik"]["ime"] or "", p["mjesec_naziv"], str(p["godina"]),
                   f(u["ukupno_sati"]), f(u["prekovremeni"]), f(u["godisnji"]), f(u["bolovanje"]), f(u["praznik"]), f(u["terenski"]),
                   f"{ozn} dan(a) s napomenom" if ozn else "")])
    w = [0.07, 0.2, 0.09, 0.06, 0.08, 0.08, 0.07, 0.07, 0.07, 0.07, 0.14]
    t = Table(pr, colWidths=[sirina * x for x in w], repeatRows=1)
    t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), svplava), ("GRID", (0, 0), (-1, -1), 0.4, C("B7C4CF")),
                           ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, traka]),
                           ("TOPPADDING", (0, 0), (-1, -1), 2.5), ("BOTTOMPADDING", (0, 0), (-1, -1), 2.5)]))
    elementi += [t]
    # --- šihtarica po radniku i mjesecu
    for p in model["sihtarice"]:
        elementi.append(PageBreak())
        r = p["radnik"]
        nasl = Table([[Paragraph(f"{p['ustanova']} | MJESEČNA EVIDENCIJA RADNOG VREMENA (ŠIHTARICA)", S["naslov"])]], colWidths=[sirina], rowHeights=[22])
        nasl.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), plava), ("VALIGN", (0, 0), (-1, -1), "MIDDLE")]))
        P = lambda txt, s="val": Paragraph(txt or "", S[s])
        zag = Table([[P("Mjesec:", "lab"), P(p["mjesec_naziv"]), P("Godina:", "lab"), P(str(p["godina"])), P("Radnik/ca:", "lab"), P(r["ime"]),
                      P("Radno mjesto:", "lab"), P(r["radno_mjesto"] or "")],
                     [P("Ugovoreno sedmično radno vrijeme (sati):", "lab"), "", "", P(f(r["ugovoreno_sedmicno"])),
                      P("Odgovorno lice:", "lab"), P(r["odgovorno_lice"] or ""), "", ""]],
                    colWidths=[sirina * x for x in (0.06, 0.1, 0.06, 0.06, 0.1, 0.25, 0.1, 0.27)])
        zag.setStyle(TableStyle([("SPAN", (0, 1), (2, 1)), ("VALIGN", (0, 0), (-1, -1), "BOTTOM"),
                                 ("LINEBELOW", (1, 0), (1, 0), 0.5, C("B7C4CF")), ("LINEBELOW", (3, 0), (3, 1), 0.5, C("B7C4CF")),
                                 ("LINEBELOW", (5, 0), (5, 1), 0.5, C("B7C4CF")), ("LINEBELOW", (7, 0), (7, 0), 0.5, C("B7C4CF")),
                                 ("TOPPADDING", (0, 0), (-1, -1), 2), ("BOTTOMPADDING", (0, 0), (-1, -1), 2)]))
        redovi = [[Paragraph(z, S["zag"]) for z in ZAGLAVLJA]]
        stilovi = [("BACKGROUND", (0, 0), (-1, 0), svplava), ("GRID", (0, 0), (-1, -1), 0.35, C("A6B4C0")),
                   ("BOX", (0, 0), (-1, -1), 0.9, C("37474F")), ("LINEBELOW", (0, 0), (-1, 0), 0.9, C("37474F")),
                   ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("TOPPADDING", (0, 0), (-1, -1), 1.2), ("BOTTOMPADDING", (0, 0), (-1, -1), 1.2),
                   ("LEFTPADDING", (0, 0), (-1, -1), 2), ("RIGHTPADDING", (0, 0), (-1, -1), 2)]
        for i, d in enumerate(p["dani"], 1):
            datum = dt.date.fromisoformat(d["datum"])
            redovi.append([Paragraph(x, S["cel"]) for x in (str(d["rb"]), datum.strftime("%d.%m.%Y"), d["dan"],
                           d["pocetak_rada"] or "", d["zavrsetak_rada"] or "", d["pocetak_pauze"] or "", d["zavrsetak_pauze"] or "",
                           f(d["ukupno_sati"]), f(d["prekovremeni"]), f(d["terenski"]), f(d["pripravnost"]), f(d["godisnji"]),
                           f(d["bolovanje"]), f(d["praznik"]), f(d["sedmicni_odmor"]))] + [Paragraph(d["napomena"] or "", S["nap"]), Paragraph(d["potpis"] or "", S["nap"])])
            if i % 2 == 0: stilovi.append(("BACKGROUND", (0, i), (-1, i), traka))
            if datum.weekday() >= 5: stilovi.append(("TEXTCOLOR", (0, i), (2, i), C("8A99A6")))
            if d["napomena"] and (d["napomena"] in E.MARKERI or any(w in d["napomena"] for w in ("nije moguće", "Nedostaje", "nisu odbijeni", "nije unesena", "nisu upisani"))):
                stilovi.append(("BACKGROUND", (15, i), (15, i), amber))
        u = p["ukupno"]
        redovi.append([Paragraph("<b>UKUPNO</b>", S["cel"]), "", "", "", "", "", ""] + [Paragraph(f"<b>{f(u[k])}</b>", S["cel"]) for k in
                      ("ukupno_sati", "prekovremeni", "terenski", "pripravnost", "godisnji", "bolovanje", "praznik", "sedmicni_odmor")] +
                      [Paragraph("<b>Pregledao/la:</b>", S["nap"]), ""])
        n = len(redovi) - 1
        stilovi += [("SPAN", (0, n), (6, n)), ("BACKGROUND", (0, n), (-1, n), zuta), ("LINEABOVE", (0, n), (-1, n), 0.9, C("37474F"))]
        tab = Table(redovi, colWidths=kol, repeatRows=1); tab.setStyle(TableStyle(stilovi))
        pot = Table([[P("Potpis radnika/ce: ______________________", "pot"), P("Potpis odgovornog lica: ______________________", "pot"),
                      P("Datum ovjere: ______________", "pot"), P("Napomene/korekcije: ______________________", "pot")]],
                    colWidths=[sirina * x for x in (0.25, 0.28, 0.2, 0.27)])
        fus = Paragraph("Napomena: Ukupno dnevno sati je automatski obračun efektivnog intervala između početka i završetka rada, umanjen za "
                        "evidentiranu pauzu. Odsustva, terenski rad i pripravnost unose se zasebno i ne dodaju se automatski u ukupno dnevno vrijeme.", S["fus"])
        elementi += [nasl, Spacer(0, 4), zag, Spacer(0, 5), tab, Spacer(0, 7), KeepTogether([pot, Spacer(0, 3), fus])]
    doc.build(elementi)
    return len(model["sihtarice"])
